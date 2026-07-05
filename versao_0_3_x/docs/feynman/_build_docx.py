"""Render docs/feynman/USEHBN-EXPLICADO.md → docs/feynman/USEHBN-EXPLICADO.docx.

Run from repo root:
    .venv/bin/python docs/feynman/_build_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt


HERE = Path(__file__).resolve().parent
SRC = HERE / "USEHBN-EXPLICADO.md"
DST = HERE / "USEHBN-EXPLICADO.docx"


def render(md_text: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    in_table = False
    table_rows: list[list[str]] = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        header = table_rows[0]
        body = table_rows[2:] if len(table_rows) > 1 else []
        tbl = doc.add_table(rows=1, cols=len(header))
        tbl.style = "Light List Accent 1"
        for i, cell in enumerate(header):
            tbl.rows[0].cells[i].text = cell.strip()
        for row in body:
            cells = tbl.add_row().cells
            for i, cell in enumerate(row):
                if i < len(cells):
                    cells[i].text = cell.strip()
        table_rows = []

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()

        if line.startswith("|") and line.endswith("|"):
            cells = [c for c in line.split("|")[1:-1]]
            table_rows.append(cells)
            in_table = True
            continue
        if in_table and not line.strip():
            flush_table()
            in_table = False
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:].strip())
            p.style = doc.styles["Intense Quote"]
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\) ", line):
            doc.add_paragraph(re.sub(r"^\d+\) ", "", line), style="List Number")
        elif line.startswith("```"):
            continue
        elif line.startswith("---"):
            doc.add_paragraph("―" * 30)
        elif not line.strip():
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)

    flush_table()
    return doc


def main(argv: list[str] | None = None) -> int:
    if not SRC.exists():
        print(f"missing source: {SRC}", file=sys.stderr)
        return 1
    doc = render(SRC.read_text(encoding="utf-8"))
    doc.save(str(DST))
    print(f"wrote {DST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
